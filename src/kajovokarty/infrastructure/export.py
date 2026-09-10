from pathlib import Path
import csv, io, json, os, tempfile, zipfile
from kajovokarty.domain.core import AppError, canonical, decimal_money, require
from kajovokarty.application.reports import SCHEMA


def columns(dataset):
    result = []
    for declaration in SCHEMA[dataset]:
        key, kind = declaration.split(":")
        result.append((key, kind))
        if key.endswith("_minor") and kind in ("I", "I?"):
            result.append((key[:-6] + "_decimal", "S?" if kind.endswith("?") else "S"))
    return result


def values(dataset, row, csv_mode=False):
    result = []
    for key, kind in columns(dataset):
        if key.endswith("_decimal"):
            n = row.get(key[:-8] + "_minor")
            v = decimal_money(n) if n is not None else None
        else:
            v = row.get(key)
        if kind.startswith("J"):
            # Persistent *_json values already have canonical serialized form.
            if isinstance(v, str) and key != "value_json":
                try:
                    v = json.loads(v)
                except ValueError:
                    raise AppError(
                        "EXPORT_INVALID", "Datová sada obsahuje neplatný JSON."
                    )
            v = canonical(v)
        elif kind.startswith("B") and v is not None:
            v = bool(v)
        if csv_mode:
            if isinstance(v, bool):
                v = "true" if v else "false"
            if (
                isinstance(v, str)
                and not key.endswith("_decimal")
                and v.lstrip().startswith(("=", "+", "-", "@"))
            ):
                v = "'" + v
        require(
            v is not None or kind.endswith("?"),
            "EXPORT_INVALID",
            "Povinný sloupec sestavy nemá hodnotu.",
            {"dataset": dataset, "column": key},
        )
        result.append(v)
    return result


def export(data, format, path, cancel=None):
    def check():
        require(not (cancel and cancel.is_set()), "CANCELLED", "Export byl zrušen.")

    check()
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    handle = tempfile.NamedTemporaryFile(
        dir=path.parent, prefix=".kk-export-", suffix="." + format, delete=False
    )
    tmp = Path(handle.name)
    handle.close()
    try:
        if format == "zip":
            with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
                for dataset, rows in data.items():
                    buf = io.StringIO(newline="")
                    writer = csv.writer(buf, lineterminator="\r\n")
                    writer.writerow([k for k, t in columns(dataset)])
                    for r in rows:
                        check()
                        writer.writerow(values(dataset, r, True))
                    z.writestr(dataset + ".csv", buf.getvalue().encode("utf-8-sig"))
            with zipfile.ZipFile(tmp) as z:
                require(
                    z.testzip() is None,
                    "EXPORT_INVALID",
                    "Exportovaný ZIP je poškozen.",
                )
        elif format == "xlsx":
            from openpyxl import Workbook
            from openpyxl.cell import WriteOnlyCell
            from openpyxl.styles import Font, Alignment

            book = Workbook(write_only=True)
            for dataset, rows in data.items():
                chunks = max(1, (len(rows) + 1048574) // 1048575)
                for index in range(chunks):
                    sh = book.create_sheet(
                        dataset if chunks == 1 else f"{dataset}_{index + 1:03d}"
                    )
                    sh.freeze_panes = "A2"
                    cols = columns(dataset)
                    heads = []
                    for key, typ in cols:
                        cell = WriteOnlyCell(sh, key)
                        cell.font = Font(bold=True)
                        heads.append(cell)
                    from openpyxl.utils import get_column_letter

                    for col_index, (key, typ) in enumerate(cols, 1):
                        sh.column_dimensions[get_column_letter(col_index)].width = (
                            42
                            if typ.startswith("J")
                            else 24
                            if typ.startswith(("S", "T"))
                            else 18
                        )
                    sh.append(heads)
                    for row in rows[index * 1048575 : (index + 1) * 1048575]:
                        check()
                        cells = []
                        for (key, typ), v in zip(cols, values(dataset, row)):
                            if (
                                typ.startswith("I")
                                and v is not None
                                and len(str(abs(v))) > 15
                            ):
                                v = str(v)
                            cell = WriteOnlyCell(sh, v)
                            if isinstance(v, str):
                                cell.data_type = "s"
                            cell.alignment = Alignment(wrap_text=True, vertical="top")
                            cells.append(cell)
                        sh.append(cells)
                    from openpyxl.utils import get_column_letter

                    sh.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{min(len(rows), 1048575) + 1}"
            book.save(tmp)
        elif format == "pdf":
            from kajovokarty.infrastructure.pdf_export import render_pdf

            render_pdf(data, tmp, check)
        else:
            raise AppError("EXPORT_INVALID", "Nepodporovaný exportní formát.")
        check()
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)
    return str(path)
